"""Tests for PDF extract, Excel extract, Gmail send options, and automated overdue reminders."""
import io
from datetime import date
from unittest.mock import patch

import openpyxl

from backend.models import Contact, Invoice


def _create_invoice(client, org, status="sent", inv_date="2026-09-01"):
    payload = {
        "customerId": org["customer"]["id"],
        "date": inv_date,
        "status": status,
        "lines": [{"itemId": org["item"]["id"], "description": "Dell Monitor 27-inch", "quantity": 1, "rate": 15000, "taxRate": 18}],
    }
    res = client.post("/api/invoices", headers=org["h"], json=payload)
    assert res.status_code == 201, res.text
    return res.json()


def test_download_invoice_pdf(client, org):
    """Verify single invoice PDF generation returns valid PDF binary data."""
    inv = _create_invoice(client, org)
    res = client.get(f"/api/invoices/{inv['id']}/pdf", headers=org["h"])
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.content.startswith(b"%PDF-")
    assert len(res.content) > 1000
    assert f"attachment; filename=\"Invoice-{inv['invoiceNumber']}.pdf\"" in res.headers["content-disposition"]


def test_download_invoice_excel(client, org):
    """Verify single invoice Excel generation returns valid openpyxl workbook."""
    inv = _create_invoice(client, org)
    res = client.get(f"/api/invoices/{inv['id']}/excel", headers=org["h"])
    assert res.status_code == 200
    assert "openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert res.content.startswith(b"PK\x03\x04")  # Standard ZIP / XLSX magic bytes

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active
    assert inv["invoiceNumber"] in ws.title or "Invoice" in ws.title
    assert "TAX INVOICE" in ws["A1"].value
    assert ws["B3"].value == inv["invoiceNumber"]


def test_export_invoices_pdf_and_excel(client, org):
    """Verify multi-invoice registry export to PDF and Excel."""
    _create_invoice(client, org)

    # PDF list export
    res_pdf = client.get("/api/invoices/export/pdf", headers=org["h"])
    assert res_pdf.status_code == 200
    assert res_pdf.headers["content-type"] == "application/pdf"
    assert res_pdf.content.startswith(b"%PDF-")

    # Excel list export
    res_excel = client.get("/api/invoices/export/excel", headers=org["h"])
    assert res_excel.status_code == 200
    assert "openxmlformats-officedocument.spreadsheetml.sheet" in res_excel.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(res_excel.content))
    ws = wb.active
    assert "Sales Invoices" in ws.title
    assert "Registry" in ws["A1"].value


def test_send_invoice_via_gmail_options(client, org):
    """Verify sending invoice via Gmail SMTP with options (Tax Invoice vs Overdue, PDF attachment)."""
    inv = _create_invoice(client, org)

    with patch("backend.routers.invoices.send_invoice_email") as mock_send_inv:
        mock_send_inv.return_value = {"success": True, "message": "Email sent", "pdf_attached": True}
        res = client.post(
            f"/api/invoices/{inv['id']}/send-gmail",
            headers=org["h"],
            json={
                "to_email": "finance@client.com",
                "send_as_overdue": False,
                "attach_pdf": True,
                "custom_notes": "Thank you for your prompt payment.",
            },
        )
        assert res.status_code == 200
        assert mock_send_inv.called
        kwargs = mock_send_inv.call_args.kwargs
        assert kwargs["to_email"] == "finance@client.com"
        assert kwargs["custom_notes"] == "Thank you for your prompt payment."
        assert kwargs["pdf_bytes"] is not None
        assert kwargs["pdf_filename"] == f"Invoice_{inv['invoiceNumber']}.pdf"

    with patch("backend.routers.invoices.send_due_reminder_email") as mock_send_remind:
        mock_send_remind.return_value = {"success": True, "message": "Reminder sent", "pdf_attached": True}
        res2 = client.post(
            f"/api/invoices/{inv['id']}/send-gmail",
            headers=org["h"],
            json={
                "to_email": "finance@client.com",
                "send_as_overdue": True,
                "attach_pdf": True,
                "custom_notes": "Urgent reminder.",
            },
        )
        assert res2.status_code == 200
        assert mock_send_remind.called
        kwargs2 = mock_send_remind.call_args.kwargs
        assert kwargs2["to_email"] == "finance@client.com"
        assert kwargs2["custom_notes"] == "Urgent reminder."


def test_auto_remind_overdue_invoices(client, org):
    """Verify scanning and automated dispatch of overdue reminders via Gmail SMTP."""
    from backend.db import SessionLocal

    with SessionLocal() as db_session:
        # Ensure customer has email
        cust = db_session.get(Contact, org["customer"]["id"])
        cust.email = "overdue.client@acme.com"
        db_session.commit()

        # Create an overdue invoice with an earlier date
        inv = _create_invoice(client, org, status="sent", inv_date="2026-07-01")

        # Force due date to past
        inv_model = db_session.get(Invoice, inv["id"])
        inv_model.due_date = date(2026, 7, 15)
        db_session.commit()

    with patch("backend.routers.invoices.send_due_reminder_email") as mock_remind:
        mock_remind.return_value = {"success": True, "message": "Reminder sent"}
        res = client.post("/api/invoices/auto-remind-overdue", headers=org["h"])
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert data["total_overdue"] >= 1
        assert data["reminders_sent"] >= 1
        assert any(d["email"] == "overdue.client@acme.com" for d in data["dispatched"])
        assert mock_remind.called
