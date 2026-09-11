import csv
import hashlib
import io
import os
import re
import uuid
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from backend.config import get_settings
from backend.db import get_db
from backend.deps import get_current_user, require_write
from backend.models import (
    Account,
    BankAccount,
    Bill,
    BillLine,
    Contact,
    Document,
    Expense,
    Invoice,
    InvoiceLine,
    User,
)
from backend.schemas.common import Message, Page
from backend.schemas.documents import DocumentOut, DocumentUpdate
from backend.services import audit, excel_import, numbering
from backend.services.tenancy import Pagination, get_or_404, paginate

settings = get_settings()
router = APIRouter(prefix="/api/documents", tags=["Documents"])

ALLOWED_TYPES = {
    "application/pdf", "image/png", "image/jpeg", "image/webp", "text/csv", "text/plain",
    "application/vnd.ms-excel", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/msword", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
CATEGORIES = {"general", "invoice", "bill", "receipt", "contract", "tax", "bank_statement", "payroll", "other"}


def to_out(doc: Document) -> DocumentOut:
    data = DocumentOut.model_validate(doc)
    data.uploaded_by_name = doc.uploader.name if doc.uploader else None
    return data


def _safe_name(name: str) -> str:
    base = os.path.basename(name or "file")
    return re.sub(r"[^A-Za-z0-9._-]+", "_", base)[:120] or "file"


@router.get("", response_model=Page[DocumentOut])
def list_documents(
    category: Optional[str] = None,
    search: Optional[str] = None,
    linked_entity_type: Optional[str] = None,
    linked_entity_id: Optional[str] = None,
    pagination: Pagination = Depends(),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    stmt = select(Document).where(Document.organization_id == user.organization_id).options(selectinload(Document.uploader))
    if category:
        stmt = stmt.where(Document.category == category)
    if linked_entity_type:
        stmt = stmt.where(Document.linked_entity_type == linked_entity_type)
    if linked_entity_id:
        stmt = stmt.where(Document.linked_entity_id == linked_entity_id)
    if search and search.strip():
        q = f"%{search.strip().lower()}%"
        stmt = stmt.where(func.lower(Document.title).like(q) | func.lower(Document.original_filename).like(q) | func.lower(Document.notes).like(q))
    stmt = stmt.order_by(Document.created_at.desc())
    rows, total = paginate(db, stmt, pagination)
    return Page(items=[to_out(d) for d in rows], total=total, page=pagination.page, page_size=pagination.page_size)


@router.post("", response_model=DocumentOut, status_code=status.HTTP_201_CREATED)
async def upload_document(
    file: UploadFile = File(...),
    title: Optional[str] = Form(None),
    category: str = Form("general"),
    notes: Optional[str] = Form(None),
    linked_entity_type: Optional[str] = Form(None),
    linked_entity_id: Optional[str] = Form(None),
    user: User = Depends(require_write),
    db: Session = Depends(get_db),
):
    if category not in CATEGORIES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Category must be one of: {', '.join(sorted(CATEGORIES))}")
    content_type = (file.content_type or "application/octet-stream").split(";")[0].strip()
    if content_type not in ALLOWED_TYPES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"File type '{content_type}' is not allowed")

    limit = settings.max_upload_size_mb * 1024 * 1024
    org_dir = os.path.join(settings.upload_dir, user.organization_id)
    os.makedirs(org_dir, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex}_{_safe_name(file.filename or 'file')}"
    stored_path = os.path.join(org_dir, stored_name)

    digest = hashlib.sha256()
    size = 0
    with open(stored_path, "wb") as fh:
        while chunk := await file.read(1024 * 1024):
            size += len(chunk)
            if size > limit:
                fh.close()
                os.remove(stored_path)
                raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, f"File exceeds the {settings.max_upload_size_mb} MB limit")
            digest.update(chunk)
            fh.write(chunk)
    if size == 0:
        os.remove(stored_path)
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Uploaded file is empty")

    doc = Document(
        organization_id=user.organization_id, title=(title or file.filename or "Document")[:255], category=category,
        original_filename=(file.filename or "file")[:255], stored_path=stored_path, content_type=content_type, size_bytes=size,
        sha256=digest.hexdigest(), notes=notes, linked_entity_type=linked_entity_type, linked_entity_id=linked_entity_id, uploaded_by=user.id,
    )
    db.add(doc)
    db.flush()
    audit.record(db, user, "create", "document", doc.id, f"Uploaded {doc.original_filename} ({size} bytes)")
    db.commit()
    db.refresh(doc)
    return to_out(doc)


# ---------------------------------------------------------------------------
# Excel / CSV Auto-Categorization & Data Input Engine
# ---------------------------------------------------------------------------
class ExcelRowIssue(BaseModel):
    row_number: int
    errors: List[str]


class ExcelCategorizeSection(BaseModel):
    category: str
    sheet_name: str
    headers: List[str]
    count: int
    rows: List[Dict[str, Any]]
    # What the preview needs to show before anything is written.
    mapped_columns: Dict[str, str] = {}
    missing_required: List[str] = []
    unmapped_headers: List[str] = []
    skipped_count: int = 0
    issues: List[ExcelRowIssue] = []
    rules: Dict[str, Any] = {}

    @property
    def importable(self) -> bool:
        return not self.missing_required and self.count > 0


class ExcelCategorizeResponse(BaseModel):
    filename: str
    total_sheets: int
    total_rows: int
    sections: List[ExcelCategorizeSection]
    # True when every sheet has the columns it needs, so the UI can block the
    # extract button rather than importing half-understood data.
    ready: bool = True
    blocking_problems: List[str] = []


class ExcelCommitItem(BaseModel):
    category: str
    data: Dict[str, Any]


class ExcelCommitPayload(BaseModel):
    items: Optional[List[ExcelCommitItem]] = None
    sections: Optional[List[ExcelCategorizeSection]] = None


class ExcelCommitResponse(BaseModel):
    success: bool
    imported_counts: Dict[str, int]
    message: str
    # Rows that could not be imported, with the reason - never silently dropped.
    skipped: List[str] = []


def _cell_to_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    return str(val).strip()


def _rows_to_section(
    sheet_name: str,
    raw_rows: List[List[str]],
    default_first_data_row: int = 2,
) -> Optional[ExcelCategorizeSection]:
    """Build one preview section from a sheet's raw rows, or None if unusable."""
    if len(raw_rows) < 2:
        return None
    # Keep blank header cells in place: dropping them shifts every later column
    # and silently files values under the wrong heading.
    headers = [h.strip() for h in raw_rows[0]]
    if not any(headers):
        return None

    category = excel_import.detect_category(sheet_name, headers)
    if category is None:
        return None

    mapping = excel_import.map_columns(headers, category, sheet_name)
    parsed = excel_import.parse_rows(raw_rows[1:], headers, mapping, first_data_row=default_first_data_row)

    good = [row for row in parsed if row.ok]
    bad = [row for row in parsed if not row.ok]

    return ExcelCategorizeSection(
        category=category,
        sheet_name=sheet_name,
        headers=headers,
        count=len(good),
        rows=[_jsonable(row.data) for row in good],
        mapped_columns=mapping.mapped,
        missing_required=mapping.missing_required,
        unmapped_headers=mapping.unmapped_headers,
        skipped_count=len(bad),
        issues=[ExcelRowIssue(row_number=row.row_number, errors=row.errors) for row in bad[:25]],
        rules=excel_import.rules_for_display(category),
    )


def _jsonable(data: Dict[str, Any]) -> Dict[str, Any]:
    """Dates/Decimals -> JSON-safe values that survive the round trip to commit."""
    out: Dict[str, Any] = {}
    for key, value in data.items():
        if isinstance(value, (date, datetime)):
            out[key] = value.isoformat()[:10]
        elif isinstance(value, Decimal):
            out[key] = float(value)
        else:
            out[key] = value
    return out


@router.post("/import-excel-categorize", response_model=ExcelCategorizeResponse)
async def import_excel_categorize(
    file: UploadFile = File(...),
    user: User = Depends(get_current_user),
):
    """Inspect an uploaded Excel/CSV file and show exactly what would be imported.

    Nothing is written here. The response carries the column mapping that was
    recognised, any required columns that are missing, and the rows that could
    not be read - so the user reviews a real preview before committing.
    """
    filename = file.filename or "import.xlsx"
    contents = await file.read()
    if not contents:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "File is empty")

    sections: List[ExcelCategorizeSection] = []

    if filename.lower().endswith((".csv", ".txt")):
        try:
            text = contents.decode("utf-8")
        except UnicodeDecodeError:
            text = contents.decode("latin-1", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        raw_rows = [[col.strip() for col in row] for row in reader if any(c.strip() for c in row)]
        section = _rows_to_section("CSV Data", raw_rows)
        if section:
            sections.append(section)
    else:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        except Exception as e:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Failed to parse Excel workbook: {e}") from e
        for sheet_name in wb.sheetnames:
            s_lower = sheet_name.strip().lower()
            if any(x in s_lower for x in ["overview", "guide", "readme", "instruction", "legend"]):
                continue
            ws = wb[sheet_name]
            raw_rows = [
                [_cell_to_str(c) for c in row_vals]
                for row_vals in ws.iter_rows(values_only=True)
                if any(v is not None and str(v).strip() for v in row_vals)
            ]
            section = _rows_to_section(sheet_name, raw_rows)
            if section:
                sections.append(section)

    if not sections:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "No importable sheets found. Name each sheet after what it holds "
            "(Customers, Vendors, Invoices, Bills or Expenses) and give it a header row.",
        )

    problems: List[str] = []
    for section in sections:
        if section.missing_required:
            problems.append(
                f"'{section.sheet_name}' is missing required column(s): {', '.join(section.missing_required)}."
            )
    return ExcelCategorizeResponse(
        filename=filename,
        total_sheets=len(sections),
        total_rows=sum(s.count for s in sections),
        sections=sections,
        ready=not problems,
        blocking_problems=problems,
    )


@router.post("/import-excel-commit", response_model=ExcelCommitResponse)
def import_excel_commit(
    payload: ExcelCommitPayload,
    user: User = Depends(require_write),
    db: Session = Depends(get_db),
):
    """Write the previewed rows into the books.

    Only values actually present in the sheet are used. A row missing a
    required value is skipped and reported rather than being filled in with an
    invented amount, date or email - importing a spreadsheet must never
    fabricate accounting data.
    """
    # Imported late: the routers import each other's services, and importing at
    # module level here would create a circular import.
    from backend.routers.bills import post_bill
    from backend.routers.invoices import post_invoice

    org_id = user.organization_id
    counts: Dict[str, int] = {"customers": 0, "vendors": 0, "invoices": 0, "bills": 0, "expenses": 0}
    skipped: List[str] = []

    expense_acct = db.execute(select(Account).where(Account.organization_id == org_id, Account.type == "expense")).scalars().first()
    bank_acct = db.execute(select(BankAccount).where(BankAccount.organization_id == org_id)).scalars().first()

    # Exactly one source of rows. Callers have sent both `sections` and an
    # `items` copy of the same rows, and appending both imported everything
    # twice. `sections` wins because only that path carries the sheet's
    # required-column check - taking `items` alongside it would also let a
    # sheet that failed validation slip in through the back door.
    all_items: List[ExcelCommitItem] = []
    if payload.sections:
        for sec in payload.sections:
            if sec.missing_required:
                skipped.append(f"Sheet '{sec.sheet_name}' skipped: missing {', '.join(sec.missing_required)}")
                continue
            for row in sec.rows:
                all_items.append(ExcelCommitItem(category=sec.category, data=row))
    elif payload.items:
        all_items.extend(payload.items)

    def _amount(d: Dict[str, Any]) -> Optional[Decimal]:
        return excel_import.coerce_amount(d.get("amount"))

    def _date_of(d: Dict[str, Any], key: str = "date") -> Optional[date]:
        return excel_import.coerce_date(d.get(key))

    def _find_or_create_contact(name: str, kind: str, email: Optional[Any]) -> Contact:
        existing = db.execute(
            select(Contact).where(Contact.organization_id == org_id, Contact.type == kind, Contact.display_name == name)
        ).scalars().first()
        if existing:
            return existing
        created = Contact(
            organization_id=org_id,
            type=kind,
            display_name=name[:100],
            email=(str(email)[:100] if email else None),
        )
        db.add(created)
        db.flush()
        return created

    for index, item in enumerate(all_items, start=1):
        cat = item.category.lower()
        d = item.data
        name = str(d.get("display_name") or "").strip()

        if cat in ("customers", "vendors"):
            if not name:
                skipped.append(f"Row {index} ({cat}): no name")
                continue
            kind = "customer" if cat == "customers" else "vendor"
            # Re-importing the same sheet must not pile up duplicates.
            already = db.execute(
                select(Contact.id).where(
                    Contact.organization_id == org_id, Contact.type == kind, Contact.display_name == name[:100]
                )
            ).first()
            if already:
                skipped.append(f"Row {index} ({cat}): '{name}' already exists")
                continue
            db.add(
                Contact(
                    organization_id=org_id,
                    type=kind,
                    display_name=name[:100],
                    company_name=str(d.get("company_name") or "")[:100] or None,
                    email=str(d.get("email") or "")[:100] or None,
                    phone=str(d.get("phone") or "")[:20] or None,
                    gstin=str(d.get("gstin") or "")[:15] or None,
                    pan=str(d.get("pan") or "")[:10] or None,
                    billing_address=str(d.get("billing_address") or "")[:255] or None,
                )
            )
            db.flush()
            counts[cat] += 1

        elif cat == "expenses":
            # No falling back to today: the sheet's Date column is required, and
            # silently dating an expense "now" files it in the wrong period.
            amount, when = _amount(d), _date_of(d)
            if amount is None or when is None:
                skipped.append(f"Row {index} (expense): missing {'amount' if amount is None else 'date'}")
                continue
            if not expense_acct or not bank_acct:
                skipped.append(f"Row {index} (expense): no expense or bank account set up yet")
                continue
            payee = str(d.get("payee") or "").strip()
            expense = Expense(
                organization_id=org_id,
                expense_number=numbering.next_number(db, org_id, "expense"),
                date=when,
                account_id=expense_acct.id,
                paid_through_account_id=bank_acct.id,
                amount=amount,
                total=amount,
                tax_rate=Decimal("0"),
                notes=str(d.get("notes") or "").strip() or None,
                created_by=user.id,
            )
            # Only set what the sheet actually provided; anything it does not
            # carry is left alone rather than filled with placeholder text.
            category = str(d.get("category") or "").strip()
            if category:
                expense.category = category[:50]
            if payee:
                expense.vendor_id = _find_or_create_contact(payee, "vendor", None).id
            db.add(expense)
            counts["expenses"] += 1

        elif cat in ("invoices", "bills"):
            amount, when = _amount(d), _date_of(d)
            if not name or amount is None or when is None:
                missing = "customer/vendor name" if not name else ("amount" if amount is None else "date")
                skipped.append(f"Row {index} ({cat[:-1]}): missing {missing}")
                continue
            # due_date is NOT NULL on the model; with no Due Date column the
            # document's own date is the only non-invented value available.
            due = _date_of(d, "due_date") or when
            # No placeholder text: if the sheet has no Notes column the line
            # description stays empty rather than saying something untrue.
            description = str(d.get("notes") or "").strip()

            if cat == "invoices":
                customer = _find_or_create_contact(name, "customer", d.get("email"))
                invoice = Invoice(
                    organization_id=org_id,
                    invoice_number=numbering.next_number(db, org_id, "invoice"),
                    customer_id=customer.id,
                    date=when,
                    due_date=due,
                    reference=str(d.get("invoice_number") or "")[:50] or None,
                    subtotal=amount,
                    tax_total=Decimal("0"),
                    total=amount,
                    amount_paid=Decimal("0"),
                    status="sent",
                    created_by=user.id,
                )
                invoice.lines.append(
                    InvoiceLine(position=0, description=description, quantity=Decimal("1"), rate=amount,
                                tax_rate=Decimal("0"), amount=amount, tax_amount=Decimal("0"))
                )
                db.add(invoice)
                db.flush()
                post_invoice(db, invoice, user)
                counts["invoices"] += 1
            else:
                vendor = _find_or_create_contact(name, "vendor", d.get("email"))
                bill = Bill(
                    organization_id=org_id,
                    bill_number=numbering.next_number(db, org_id, "bill"),
                    vendor_bill_number=str(d.get("bill_number") or "")[:50] or None,
                    vendor_id=vendor.id,
                    date=when,
                    due_date=due,
                    subtotal=amount,
                    tax_total=Decimal("0"),
                    total=amount,
                    amount_paid=Decimal("0"),
                    status="open",
                    created_by=user.id,
                )
                bill.lines.append(
                    BillLine(position=0, description=description, quantity=Decimal("1"), rate=amount,
                             tax_rate=Decimal("0"), amount=amount, tax_amount=Decimal("0"))
                )
                db.add(bill)
                db.flush()
                post_bill(db, bill, user)
                counts["bills"] += 1

    db.commit()
    total_saved = sum(counts.values())
    audit.record(db, user, "import", "excel_input", user.id, f"Imported {total_saved} records: {counts}")
    db.commit()

    message = f"Imported {total_saved} record(s)."
    if skipped:
        message += f" {len(skipped)} row(s) were skipped."
    return ExcelCommitResponse(
        success=True,
        imported_counts=counts,
        message=message,
        skipped=skipped[:50],
    )


@router.get("/download-sample-excel")
def download_sample_excel(type: str = "indian"):
    filename = "Rooman_Books_Indian_Data.xlsx" if type == "indian" else "Rooman_Books_Import_Template.xlsx"
    repo_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    file_path = os.path.join(repo_root, "frontend", "public", filename)
    if not os.path.exists(file_path):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Sample Excel file not found")
    return FileResponse(
        file_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-cache",
        },
    )


@router.get("/{document_id}", response_model=DocumentOut)
def get_document(document_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return to_out(get_or_404(db, Document, document_id, user.organization_id, "Document"))


@router.get("/{document_id}/download")
def download_document(document_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    doc = get_or_404(db, Document, document_id, user.organization_id, "Document")
    if not os.path.exists(doc.stored_path):
        raise HTTPException(status.HTTP_410_GONE, "Stored file is missing")
    return FileResponse(doc.stored_path, media_type=doc.content_type, filename=doc.original_filename)


@router.patch("/{document_id}", response_model=DocumentOut)
def update_document(document_id: str, payload: DocumentUpdate, user: User = Depends(require_write), db: Session = Depends(get_db)):
    doc = get_or_404(db, Document, document_id, user.organization_id, "Document")
    data = payload.model_dump(exclude_unset=True)
    if "category" in data and data["category"] not in CATEGORIES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Invalid category")
    for field, value in data.items():
        setattr(doc, field, value)
    db.commit()
    db.refresh(doc)
    return to_out(doc)


@router.delete("/{document_id}", response_model=Message)
def delete_document(document_id: str, user: User = Depends(require_write), db: Session = Depends(get_db)):
    doc = get_or_404(db, Document, document_id, user.organization_id, "Document")
    path = doc.stored_path
    name = doc.original_filename
    db.delete(doc)
    audit.record(db, user, "delete", "document", document_id, f"Deleted {name}")
    db.commit()
    try:
        if os.path.exists(path):
            os.remove(path)
    except OSError:  # pragma: no cover
        pass
    return Message(message="Document deleted")



