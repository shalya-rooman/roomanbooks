"""Service for generating professional PDF and Excel extracts for Invoices and Sales Registries."""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from backend.models import Bill, Contact, CustomerPayment, Expense, Invoice, Organization, VendorPayment


def _fmt_curr(val: Optional[float]) -> str:
    amount = float(val or 0.0)
    return f"INR {amount:,.2f}"


def _fmt_date(d: Optional[date | datetime]) -> str:
    if not d:
        return "-"
    if isinstance(d, datetime):
        return d.strftime("%d %b %Y")
    return d.strftime("%d %b %Y")


def generate_invoice_pdf(invoice: Invoice, org: Optional[Organization] = None) -> bytes:
    """Generate a high-quality, professional GST Tax Invoice PDF document."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "InvoiceTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0f172a"),
    )
    subtitle_style = ParagraphStyle(
        "InvoiceSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#2563eb"),
    )
    bold_label = ParagraphStyle(
        "BoldLabel",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#475569"),
    )
    value_style = ParagraphStyle(
        "ValueStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#0f172a"),
    )
    header_th = ParagraphStyle(
        "THStyle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        textColor=colors.white,
    )
    cell_style = ParagraphStyle(
        "CellStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1e293b"),
    )
    cell_right = ParagraphStyle(
        "CellRight",
        parent=cell_style,
        alignment=2,  # Right aligned
    )

    story = []

    # Org Info & Header
    org_name = (org.name if org else None) or "Rooman Technologies Pvt Ltd"
    org_addr = (org.address if org else None) or "Rooman House, #12 Rajajinagar"
    org_city_state = ", ".join(filter(None, [org.city if org else "Bengaluru", org.state if org else "Karnataka", org.postal_code if org else "560010"]))
    org_gstin = (org.gstin if org else None) or "29AABCR1234F1Z5"
    org_email = (org.email if org else None) or "shalya@rooman.com"

    header_table_data = [
        [
            Paragraph(f"<b>{org_name}</b>", title_style),
            Paragraph("<b>TAX INVOICE</b>", subtitle_style),
        ],
        [
            Paragraph(f"{org_addr}<br/>{org_city_state}<br/>GSTIN: {org_gstin}<br/>Email: {org_email}", value_style),
            Paragraph(
                f"<b>Invoice #:</b> {invoice.invoice_number}<br/>"
                f"<b>Date:</b> {_fmt_date(invoice.date)}<br/>"
                f"<b>Due Date:</b> {_fmt_date(invoice.due_date)}<br/>"
                f"<b>Status:</b> {invoice.status.upper()}",
                value_style,
            ),
        ],
    ]
    t_header = Table(header_table_data, colWidths=[330, 210])
    t_header.setStyle(
        TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ])
    )
    story.append(t_header)
    story.append(Spacer(1, 14))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=14))

    # Bill To Box
    cust = invoice.customer
    cust_name = cust.display_name if cust else "Client"
    cust_addr = (cust.billing_address if cust else "") or "Address on file"
    cust_email = (cust.email if cust else "") or "-"
    cust_gstin = (cust.gstin if cust else "") or "-"

    bill_to_data = [
        [
            Paragraph("<b>BILLED TO:</b>", bold_label),
            Paragraph("<b>PAYMENT TERMS & REFERENCE:</b>", bold_label),
        ],
        [
            Paragraph(f"<b>{cust_name}</b><br/>{cust_addr}<br/>Email: {cust_email}<br/>GSTIN: {cust_gstin}", value_style),
            Paragraph(
                f"Reference: {invoice.reference or 'N/A'}<br/>"
                f"Payment Mode: Bank / UPI<br/>"
                f"Bank: HDFC Bank - A/C 50200012345678<br/>"
                f"IFSC: HDFC0001234",
                value_style,
            ),
        ],
    ]
    t_bill = Table(bill_to_data, colWidths=[330, 210])
    t_bill.setStyle(
        TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ])
    )
    story.append(t_bill)
    story.append(Spacer(1, 16))

    # Line Items Table
    lines_header = [
        Paragraph("#", header_th),
        Paragraph("Item / Description", header_th),
        Paragraph("Qty", header_th),
        Paragraph("Rate", header_th),
        Paragraph("Tax %", header_th),
        Paragraph("Amount", header_th),
    ]
    table_rows = [lines_header]

    for idx, line in enumerate(invoice.lines, start=1):
        desc = line.description or (line.item.name if line.item else "Product / Service")
        table_rows.append([
            Paragraph(str(idx), cell_style),
            Paragraph(desc, cell_style),
            Paragraph(f"{float(line.quantity):g}", cell_right),
            Paragraph(_fmt_curr(line.rate), cell_right),
            Paragraph(f"{float(line.tax_rate or 0):g}%", cell_right),
            Paragraph(_fmt_curr(line.amount), cell_right),
        ])

    t_lines = Table(table_rows, colWidths=[24, 256, 50, 70, 50, 90])
    t_lines.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
            ("ALIGN", (0, 0), (0, -1), "CENTER"),
            ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ])
    )
    story.append(t_lines)
    story.append(Spacer(1, 14))

    # Totals Summary
    subtotal = float(invoice.subtotal or 0.0)
    tax_total = float(invoice.tax_total or 0.0)
    discount = float(invoice.discount_amount or 0.0)
    total = float(invoice.total or 0.0)
    paid = float(invoice.amount_paid or 0.0)
    balance = float(invoice.balance_due or 0.0)

    summary_rows = [
        [Paragraph("Subtotal:", cell_right), Paragraph(_fmt_curr(subtotal), cell_right)],
        [Paragraph("Taxes (GST):", cell_right), Paragraph(_fmt_curr(tax_total), cell_right)],
    ]
    if discount > 0:
        summary_rows.append([Paragraph("Discount:", cell_right), Paragraph(f"- {_fmt_curr(discount)}", cell_right)])
    summary_rows.extend([
        [Paragraph("<b>Grand Total:</b>", cell_right), Paragraph(f"<b>{_fmt_curr(total)}</b>", cell_right)],
        [Paragraph("Amount Paid:", cell_right), Paragraph(_fmt_curr(paid), cell_right)],
        [
            Paragraph("<b>Balance Due:</b>", cell_right),
            Paragraph(f"<font color='#b91c1c'><b>{_fmt_curr(balance)}</b></font>", cell_right),
        ],
    ])

    t_summary = Table(summary_rows, colWidths=[120, 100])
    t_summary.setStyle(
        TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LINEBELOW", (0, 2), (-1, 2), 0.5, colors.HexColor("#cbd5e1")),
            ("LINEBELOW", (0, -2), (-1, -2), 0.5, colors.HexColor("#cbd5e1")),
        ])
    )

    t_wrap = Table([[Paragraph(f"<b>Notes:</b><br/>{invoice.notes or 'Thank you for your business.'}", value_style), t_summary]], colWidths=[320, 220])
    t_wrap.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(t_wrap)

    story.append(Spacer(1, 25))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cbd5e1"), spaceAfter=10))
    story.append(Paragraph(
        "<i>This is a computer-generated GST Tax Invoice issued by Rooman Technologies Pvt Ltd. For inquiries, email shalya@rooman.com</i>",
        ParagraphStyle("Footnote", parent=styles["Normal"], fontSize=7.5, leading=10, textColor=colors.HexColor("#64748b"), alignment=1),
    ))

    doc.build(story)
    return buffer.getvalue()


def generate_invoices_list_pdf(invoices: List[Invoice], org: Optional[Organization] = None) -> bytes:
    """Generate a multi-row Sales Invoice Registry PDF report."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=24, leftMargin=24, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()

    th_style = ParagraphStyle("TH", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8, textColor=colors.white)
    td_style = ParagraphStyle("TD", parent=styles["Normal"], fontName="Helvetica", fontSize=7.5, leading=9)
    td_right = ParagraphStyle("TDR", parent=td_style, alignment=2)

    story = [
        Paragraph(f"<b>{(org.name if org else None) or 'Rooman Books'} - Invoices Summary Report</b>", styles["Heading2"]),
        Paragraph(f"Generated on {datetime.now().strftime('%d %b %Y, %I:%M %p')} · Total Invoices: {len(invoices)}", styles["Normal"]),
        Spacer(1, 12),
    ]

    header = [
        Paragraph("Invoice #", th_style),
        Paragraph("Customer", th_style),
        Paragraph("Date", th_style),
        Paragraph("Due Date", th_style),
        Paragraph("Status", th_style),
        Paragraph("Total", th_style),
        Paragraph("Paid", th_style),
        Paragraph("Balance Due", th_style),
    ]
    rows = [header]

    total_sum = 0.0
    paid_sum = 0.0
    balance_sum = 0.0

    for inv in invoices:
        tot = float(inv.total or 0.0)
        pd = float(inv.amount_paid or 0.0)
        bal = float(inv.balance_due or 0.0)
        total_sum += tot
        paid_sum += pd
        balance_sum += bal

        rows.append([
            Paragraph(inv.invoice_number, td_style),
            Paragraph((inv.customer.display_name if inv.customer else "-")[:24], td_style),
            Paragraph(_fmt_date(inv.date), td_style),
            Paragraph(_fmt_date(inv.due_date), td_style),
            Paragraph(inv.status.upper(), td_style),
            Paragraph(_fmt_curr(tot), td_right),
            Paragraph(_fmt_curr(pd), td_right),
            Paragraph(_fmt_curr(bal), td_right),
        ])

    rows.append([
        Paragraph("<b>TOTALS</b>", td_style),
        Paragraph("", td_style),
        Paragraph("", td_style),
        Paragraph("", td_style),
        Paragraph("", td_style),
        Paragraph(f"<b>{_fmt_curr(total_sum)}</b>", td_right),
        Paragraph(f"<b>{_fmt_curr(paid_sum)}</b>", td_right),
        Paragraph(f"<b>{_fmt_curr(balance_sum)}</b>", td_right),
    ])

    t = Table(rows, colWidths=[75, 125, 55, 55, 55, 65, 65, 69])
    t.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f8fafc")]),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e2e8f0")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ])
    )
    story.append(t)
    doc.build(story)
    return buffer.getvalue()


def generate_invoice_excel(invoice: Invoice, org: Optional[Organization] = None) -> bytes:
    """Generate a clean, structured .xlsx Excel workbook for a single invoice."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"Invoice {invoice.invoice_number}"

    # Styles
    title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    org_name = (org.name if org else None) or "Rooman Technologies Pvt Ltd"

    # Title
    ws.merge_cells("A1:F1")
    ws["A1"] = f"{org_name} - TAX INVOICE"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    # Meta
    ws["A3"] = "Invoice Number:"
    ws["B3"] = invoice.invoice_number
    ws["D3"] = "Invoice Date:"
    ws["E3"] = _fmt_date(invoice.date)

    ws["A4"] = "Customer Name:"
    ws["B4"] = invoice.customer.display_name if invoice.customer else "-"
    ws["D4"] = "Due Date:"
    ws["E4"] = _fmt_date(invoice.due_date)

    ws["A5"] = "Customer Email:"
    ws["B5"] = invoice.customer.email if invoice.customer else "-"
    ws["D5"] = "Status:"
    ws["E5"] = invoice.status.upper()

    for r in range(3, 6):
        ws[f"A{r}"].font = bold_font
        ws[f"D{r}"].font = bold_font
        ws[f"B{r}"].font = regular_font
        ws[f"E{r}"].font = regular_font

    # Items Header
    row_idx = 7
    headers = ["#", "Item / Description", "Quantity", "Unit Rate (INR)", "Tax %", "Line Amount (INR)"]
    for col_idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="right" if col_idx >= 3 else "left", vertical="center")

    # Items Rows
    for idx, line in enumerate(invoice.lines, start=1):
        row_idx += 1
        desc = line.description or (line.item.name if line.item else "Service")
        ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=desc)
        ws.cell(row=row_idx, column=3, value=float(line.quantity)).number_format = "#,##0.00"
        ws.cell(row=row_idx, column=4, value=float(line.rate)).number_format = "₹#,##0.00"
        ws.cell(row=row_idx, column=5, value=float(line.tax_rate or 0)).number_format = "0.0%"
        ws.cell(row=row_idx, column=6, value=float(line.amount)).number_format = "₹#,##0.00"

        for col_idx in range(1, 7):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = regular_font
            cell.border = thin_border

    # Totals Section
    totals_data = [
        ("Subtotal:", float(invoice.subtotal or 0.0)),
        ("GST / Tax:", float(invoice.tax_total or 0.0)),
        ("Grand Total:", float(invoice.total or 0.0)),
        ("Amount Paid:", float(invoice.amount_paid or 0.0)),
        ("Balance Due:", float(invoice.balance_due or 0.0)),
    ]

    for label, val in totals_data:
        row_idx += 1
        ws.cell(row=row_idx, column=5, value=label).font = bold_font
        ws.cell(row=row_idx, column=5).alignment = Alignment(horizontal="right")
        ws.cell(row=row_idx, column=5).fill = sub_fill

        c = ws.cell(row=row_idx, column=6, value=val)
        c.font = bold_font
        c.number_format = "₹#,##0.00"
        c.fill = sub_fill
        c.border = thin_border

    # Adjust Column Widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_invoices_list_excel(invoices: List[Invoice], org: Optional[Organization] = None) -> bytes:
    """Generate a full .xlsx Excel spreadsheet export of invoices."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Invoices"

    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    org_name = (org.name if org else None) or "Rooman Books"

    ws["A1"] = f"{org_name} - Sales Invoices Registry"
    ws["A1"].font = title_font
    ws["A2"] = f"Exported on: {datetime.now().strftime('%d %b %Y, %I:%M %p')}"
    ws["A2"].font = regular_font

    headers = [
        "Invoice Number",
        "Date",
        "Due Date",
        "Customer Name",
        "Customer Email",
        "Status",
        "Subtotal",
        "Tax Total",
        "Grand Total",
        "Amount Paid",
        "Balance Due",
    ]

    row_idx = 4
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="right" if col_idx >= 7 else "left", vertical="center")

    for inv in invoices:
        row_idx += 1
        ws.cell(row=row_idx, column=1, value=inv.invoice_number)
        ws.cell(row=row_idx, column=2, value=_fmt_date(inv.date))
        ws.cell(row=row_idx, column=3, value=_fmt_date(inv.due_date))
        ws.cell(row=row_idx, column=4, value=inv.customer.display_name if inv.customer else "-")
        ws.cell(row=row_idx, column=5, value=inv.customer.email if inv.customer else "-")
        ws.cell(row=row_idx, column=6, value=inv.status.upper())

        ws.cell(row=row_idx, column=7, value=float(inv.subtotal or 0.0)).number_format = "₹#,##0.00"
        ws.cell(row=row_idx, column=8, value=float(inv.tax_total or 0.0)).number_format = "₹#,##0.00"
        ws.cell(row=row_idx, column=9, value=float(inv.total or 0.0)).number_format = "₹#,##0.00"
        ws.cell(row=row_idx, column=10, value=float(inv.amount_paid or 0.0)).number_format = "₹#,##0.00"
        ws.cell(row=row_idx, column=11, value=float(inv.balance_due or 0.0)).number_format = "₹#,##0.00"

        for col_idx in range(1, 12):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = regular_font
            c.border = border

    # Totals row
    row_idx += 1
    ws.cell(row=row_idx, column=1, value="TOTALS").font = bold_font
    for c_idx in range(7, 12):
        col_letter = get_column_letter(c_idx)
        cell = ws.cell(row=row_idx, column=c_idx, value=f"=SUM({col_letter}5:{col_letter}{row_idx-1})")
        cell.font = bold_font
        cell.number_format = "₹#,##0.00"
        cell.border = border

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Bills & Purchases PDF / Excel
# ---------------------------------------------------------------------------


def generate_bill_pdf(bill: Bill, org: Optional[Organization] = None) -> bytes:
    """Generate a clean Purchase Bill PDF document."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle("BTitle", parent=styles["Heading1"], fontSize=18, leading=22, textColor=colors.HexColor("#0f172a"))
    sub_style = ParagraphStyle("BSub", parent=styles["Normal"], fontSize=12, fontName="Helvetica-Bold", textColor=colors.HexColor("#dc2626"))
    val_style = ParagraphStyle("BVal", parent=styles["Normal"], fontSize=9, leading=12)
    th_style = ParagraphStyle("BTH", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9, textColor=colors.white)
    cell_style = ParagraphStyle("BCell", parent=styles["Normal"], fontSize=8.5, leading=11)
    cell_right = ParagraphStyle("BCellR", parent=cell_style, alignment=2)

    org_name = (org.name if org else None) or "Rooman Technologies Pvt Ltd"
    vendor_name = bill.vendor.display_name if bill.vendor else "Vendor"
    vendor_email = (bill.vendor.email if bill.vendor else "") or "-"
    vendor_gstin = (bill.vendor.gstin if bill.vendor else "") or "-"

    story = [
        Table([
            [Paragraph(f"<b>{org_name}</b>", title_style), Paragraph("<b>PURCHASE BILL VOUCHER</b>", sub_style)],
            [
                Paragraph(f"Accounts Payable Dept<br/>GSTIN: {(org.gstin if org else None) or '29AABCR1234F1Z5'}", val_style),
                Paragraph(f"<b>Bill #:</b> {bill.bill_number}<br/><b>Vendor Ref:</b> {getattr(bill, 'vendor_bill_number', None) or 'N/A'}<br/><b>Date:</b> {_fmt_date(bill.date)}<br/><b>Due:</b> {_fmt_date(bill.due_date)}<br/><b>Status:</b> {bill.status.upper()}", val_style),
            ],
        ], colWidths=[330, 210], style=[("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (1, 0), (1, -1), "RIGHT")]),
        Spacer(1, 14),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=14),
        Table([
            [Paragraph("<b>VENDOR DETAILS:</b>", val_style), Paragraph("<b>PAYMENT STATUS:</b>", val_style)],
            [
                Paragraph(f"<b>{vendor_name}</b><br/>Email: {vendor_email}<br/>GSTIN: {vendor_gstin}", val_style),
                Paragraph(f"Total Bill Amount: <b>{_fmt_curr(bill.total)}</b><br/>Paid to date: {_fmt_curr(bill.amount_paid)}<br/>Balance Outstanding: <font color='#dc2626'><b>{_fmt_curr(bill.balance_due)}</b></font>", val_style),
            ],
        ], colWidths=[330, 210], style=[("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")), ("PADDING", (0, 0), (-1, -1), 8)]),
        Spacer(1, 16),
    ]

    # Lines
    rows = [[
        Paragraph("#", th_style),
        Paragraph("Item / Expense Description", th_style),
        Paragraph("Qty", th_style),
        Paragraph("Rate", th_style),
        Paragraph("Tax %", th_style),
        Paragraph("Amount", th_style),
    ]]
    for idx, line in enumerate(bill.lines, start=1):
        desc = line.description or (line.item.name if line.item else "Purchase Item")
        rows.append([
            Paragraph(str(idx), cell_style),
            Paragraph(desc, cell_style),
            Paragraph(f"{float(line.quantity):g}", cell_right),
            Paragraph(_fmt_curr(line.rate), cell_right),
            Paragraph(f"{float(line.tax_rate or 0):g}%", cell_right),
            Paragraph(_fmt_curr(line.amount), cell_right),
        ])

    t_lines = Table(rows, colWidths=[24, 256, 50, 70, 50, 90])
    t_lines.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(t_lines)
    story.append(Spacer(1, 14))

    # Summary
    summary = Table([
        [Paragraph("Subtotal:", cell_right), Paragraph(_fmt_curr(bill.subtotal), cell_right)],
        [Paragraph("Tax (GST):", cell_right), Paragraph(_fmt_curr(bill.tax_total), cell_right)],
        [Paragraph("<b>Total:</b>", cell_right), Paragraph(f"<b>{_fmt_curr(bill.total)}</b>", cell_right)],
        [Paragraph("Balance Due:", cell_right), Paragraph(f"<font color='#dc2626'><b>{_fmt_curr(bill.balance_due)}</b></font>", cell_right)],
    ], colWidths=[120, 100])
    summary.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "RIGHT"), ("PADDING", (0, 0), (-1, -1), 3)]))
    story.append(Table([[Paragraph(f"Notes: {bill.notes or 'Purchase bill recorded in Rooman Books.'}", val_style), summary]], colWidths=[320, 220]))

    doc.build(story)
    return buffer.getvalue()


def generate_bills_list_pdf(bills: List[Bill], org: Optional[Organization] = None) -> bytes:
    """Generate a Purchases Bill Registry PDF report."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=24, leftMargin=24, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    th_style = ParagraphStyle("BTH", fontName="Helvetica-Bold", fontSize=8, textColor=colors.white)
    td_style = ParagraphStyle("BTD", fontName="Helvetica", fontSize=7.5, leading=9)
    td_right = ParagraphStyle("BTDR", parent=td_style, alignment=2)

    story = [
        Paragraph(f"<b>{(org.name if org else None) or 'Rooman Books'} - Purchase Bills Registry</b>", styles["Heading2"]),
        Paragraph(f"Generated on {datetime.now().strftime('%d %b %Y, %I:%M %p')} · Total Bills: {len(bills)}", styles["Normal"]),
        Spacer(1, 12),
    ]

    rows = [[
        Paragraph("Bill #", th_style),
        Paragraph("Vendor", th_style),
        Paragraph("Date", th_style),
        Paragraph("Due Date", th_style),
        Paragraph("Status", th_style),
        Paragraph("Total", th_style),
        Paragraph("Paid", th_style),
        Paragraph("Balance Due", th_style),
    ]]
    for b in bills:
        rows.append([
            Paragraph(b.bill_number, td_style),
            Paragraph((b.vendor.display_name if b.vendor else "-")[:24], td_style),
            Paragraph(_fmt_date(b.date), td_style),
            Paragraph(_fmt_date(b.due_date), td_style),
            Paragraph(b.status.upper(), td_style),
            Paragraph(_fmt_curr(b.total), td_right),
            Paragraph(_fmt_curr(b.amount_paid), td_right),
            Paragraph(_fmt_curr(b.balance_due), td_right),
        ])

    t = Table(rows, colWidths=[75, 125, 55, 55, 55, 65, 65, 69])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
    doc.build(story)
    return buffer.getvalue()


def generate_bills_list_excel(bills: List[Bill], org: Optional[Organization] = None) -> bytes:
    """Export Purchase Bills to Excel workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase Bills"
    ws["A1"] = f"{(org.name if org else None) or 'Rooman Books'} - Purchase Bills Registry"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    ws["A2"] = f"Exported: {datetime.now().strftime('%d %b %Y, %I:%M %p')}"

    headers = ["Bill Number", "Vendor Name", "Vendor Email", "Date", "Due Date", "Status", "Subtotal", "Tax Total", "Total", "Paid", "Balance Due"]
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c_idx, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    for r_idx, b in enumerate(bills, start=5):
        ws.cell(row=r_idx, column=1, value=b.bill_number)
        ws.cell(row=r_idx, column=2, value=b.vendor.display_name if b.vendor else "-")
        ws.cell(row=r_idx, column=3, value=b.vendor.email if b.vendor else "-")
        ws.cell(row=r_idx, column=4, value=_fmt_date(b.date))
        ws.cell(row=r_idx, column=5, value=_fmt_date(b.due_date))
        ws.cell(row=r_idx, column=6, value=b.status.upper())
        ws.cell(row=r_idx, column=7, value=float(b.subtotal or 0.0)).number_format = "₹#,##0.00"
        ws.cell(row=r_idx, column=8, value=float(b.tax_total or 0.0)).number_format = "₹#,##0.00"
        ws.cell(row=r_idx, column=9, value=float(b.total or 0.0)).number_format = "₹#,##0.00"
        ws.cell(row=r_idx, column=10, value=float(b.amount_paid or 0.0)).number_format = "₹#,##0.00"
        ws.cell(row=r_idx, column=11, value=float(b.balance_due or 0.0)).number_format = "₹#,##0.00"

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Customer & Vendor Payment Vouchers PDF / Excel
# ---------------------------------------------------------------------------


def generate_customer_payment_pdf(payment: CustomerPayment, org: Optional[Organization] = None) -> bytes:
    """Generate official Customer Payment Acknowledgment Receipt PDF."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    org_name = (org.name if org else None) or "Rooman Technologies Pvt Ltd"
    cust_name = payment.customer.display_name if payment.customer else "Valued Customer"

    story = [
        Paragraph(f"<b>{org_name}</b>", ParagraphStyle("PTitle", parent=styles["Heading1"], fontSize=18, textColor=colors.HexColor("#0f172a"))),
        Paragraph("<b>OFFICIAL PAYMENT RECEIPT</b>", ParagraphStyle("PSub", parent=styles["Normal"], fontSize=12, fontName="Helvetica-Bold", textColor=colors.HexColor("#16a34a"))),
        Spacer(1, 10),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=14),
        Table([
            [Paragraph("<b>RECEIPT DETAILS:</b>", styles["Normal"]), Paragraph("<b>RECEIVED FROM:</b>", styles["Normal"])],
            [
                Paragraph(
                    f"Receipt #: <b>{payment.payment_number}</b><br/>"
                    f"Payment Date: {_fmt_date(payment.date)}<br/>"
                    f"Payment Mode: {payment.mode.upper()}<br/>"
                    f"Reference / UTR: {payment.reference or 'N/A'}<br/>"
                    f"Deposited To: {payment.bank_account.name if payment.bank_account else 'Bank'}",
                    styles["Normal"],
                ),
                Paragraph(
                    f"<b>{cust_name}</b><br/>"
                    f"Email: {(payment.customer.email if payment.customer else '') or '-'}<br/>"
                    f"Invoice Reference: {payment.invoice.invoice_number if payment.invoice else 'General Advance'}",
                    styles["Normal"],
                ),
            ],
        ], colWidths=[270, 270], style=[("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")), ("PADDING", (0, 0), (-1, -1), 10)]),
        Spacer(1, 20),
        Table([
            [Paragraph("<b>AMOUNT RECEIVED:</b>", ParagraphStyle("AR", fontSize=14, fontName="Helvetica-Bold")), Paragraph(f"<font color='#16a34a'><b>{_fmt_curr(payment.amount)}</b></font>", ParagraphStyle("ARVal", fontSize=18, fontName="Helvetica-Bold", alignment=2))],
        ], colWidths=[270, 270], style=[("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#ecfdf5")), ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#6ee7b7")), ("PADDING", (0, 0), (-1, -1), 12)]),
        Spacer(1, 20),
        Paragraph(f"<b>Notes:</b> {payment.notes or 'Payment credited and posted to general ledger accounts.'}", styles["Normal"]),
        Spacer(1, 30),
        Paragraph("<i>Thank you for your business. Computer generated receipt issued by Rooman Books.</i>", ParagraphStyle("Foot", fontSize=8, alignment=1, textColor=colors.HexColor("#64748b"))),
    ]
    doc.build(story)
    return buffer.getvalue()


def generate_customer_payments_list_excel(payments: List[CustomerPayment], org: Optional[Organization] = None) -> bytes:
    """Export Customer Payments received to Excel workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Payments Received"
    ws["A1"] = f"{(org.name if org else None) or 'Rooman Books'} - Payments Received Ledger"
    ws["A2"] = f"Exported: {datetime.now().strftime('%d %b %Y, %I:%M %p')}"

    headers = ["Payment #", "Date", "Customer", "Mode", "Reference", "Invoice #", "Bank Account", "Amount (INR)"]
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c_idx, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")

    for r_idx, p in enumerate(payments, start=5):
        ws.cell(row=r_idx, column=1, value=p.payment_number)
        ws.cell(row=r_idx, column=2, value=_fmt_date(p.date))
        ws.cell(row=r_idx, column=3, value=p.customer.display_name if p.customer else "-")
        ws.cell(row=r_idx, column=4, value=p.mode.upper())
        ws.cell(row=r_idx, column=5, value=p.reference or "-")
        ws.cell(row=r_idx, column=6, value=p.invoice.invoice_number if p.invoice else "-")
        ws.cell(row=r_idx, column=7, value=p.bank_account.name if p.bank_account else "-")
        ws.cell(row=r_idx, column=8, value=float(p.amount or 0.0)).number_format = "₹#,##0.00"

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_customer_payments_list_pdf(payments: List[CustomerPayment], org: Optional[Organization] = None) -> bytes:
    """Generate Customer Payments Received Registry PDF."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=24, leftMargin=24, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    th_style = ParagraphStyle("CPTH", fontName="Helvetica-Bold", fontSize=8, textColor=colors.white)
    td_style = ParagraphStyle("CPTD", fontName="Helvetica", fontSize=7.5, leading=9)
    td_right = ParagraphStyle("CPTDR", parent=td_style, alignment=2)

    story = [
        Paragraph(f"<b>{(org.name if org else None) or 'Rooman Books'} - Payments Received Registry</b>", styles["Heading2"]),
        Paragraph(f"Generated on {datetime.now().strftime('%d %b %Y, %I:%M %p')} · Total Payments: {len(payments)}", styles["Normal"]),
        Spacer(1, 12),
    ]

    rows = [[
        Paragraph("Receipt #", th_style),
        Paragraph("Customer", th_style),
        Paragraph("Date", th_style),
        Paragraph("Mode", th_style),
        Paragraph("Reference", th_style),
        Paragraph("Invoice #", th_style),
        Paragraph("Amount", th_style),
    ]]
    for p in payments:
        rows.append([
            Paragraph(p.payment_number, td_style),
            Paragraph((p.customer.display_name if p.customer else "-")[:24], td_style),
            Paragraph(_fmt_date(p.date), td_style),
            Paragraph(p.mode.upper(), td_style),
            Paragraph(p.reference or "-", td_style),
            Paragraph(p.invoice.invoice_number if p.invoice else "Advance", td_style),
            Paragraph(_fmt_curr(p.amount), td_right),
        ])

    t = Table(rows, colWidths=[80, 130, 60, 60, 85, 75, 70])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16a34a")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
    doc.build(story)
    return buffer.getvalue()


def generate_vendor_payment_pdf(payment: VendorPayment, org: Optional[Organization] = None) -> bytes:
    """Generate official Vendor Payment Voucher PDF."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    org_name = (org.name if org else None) or "Rooman Technologies Pvt Ltd"
    vendor_name = payment.vendor.display_name if payment.vendor else "Vendor"

    story = [
        Paragraph(f"<b>{org_name}</b>", ParagraphStyle("VPTitle", parent=styles["Heading1"], fontSize=18, textColor=colors.HexColor("#0f172a"))),
        Paragraph("<b>PAYMENT REMITTANCE ADVICE</b>", ParagraphStyle("VPSub", parent=styles["Normal"], fontSize=12, fontName="Helvetica-Bold", textColor=colors.HexColor("#2563eb"))),
        Spacer(1, 10),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=14),
        Table([
            [Paragraph("<b>VOUCHER DETAILS:</b>", styles["Normal"]), Paragraph("<b>PAID TO:</b>", styles["Normal"])],
            [
                Paragraph(
                    f"Voucher #: <b>{payment.payment_number}</b><br/>"
                    f"Date: {_fmt_date(payment.date)}<br/>"
                    f"Mode: {payment.mode.upper()}<br/>"
                    f"Ref / Cheque #: {payment.reference or 'N/A'}<br/>"
                    f"Paid From: {payment.bank_account.name if payment.bank_account else 'Bank'}",
                    styles["Normal"],
                ),
                Paragraph(
                    f"<b>{vendor_name}</b><br/>"
                    f"Email: {(payment.vendor.email if payment.vendor else '') or '-'}<br/>"
                    f"Bill Applied: {payment.bill.bill_number if payment.bill else 'Direct Vendor Settlement'}",
                    styles["Normal"],
                ),
            ],
        ], colWidths=[270, 270], style=[("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")), ("PADDING", (0, 0), (-1, -1), 10)]),
        Spacer(1, 20),
        Table([
            [Paragraph("<b>AMOUNT DISBURSED:</b>", ParagraphStyle("AD", fontSize=14, fontName="Helvetica-Bold")), Paragraph(f"<font color='#2563eb'><b>{_fmt_curr(payment.amount)}</b></font>", ParagraphStyle("ADVal", fontSize=18, fontName="Helvetica-Bold", alignment=2))],
        ], colWidths=[270, 270], style=[("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#eff6ff")), ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#93c5fd")), ("PADDING", (0, 0), (-1, -1), 12)]),
        Spacer(1, 20),
        Paragraph(f"<b>Notes:</b> {payment.notes or 'Vendor payment processed and confirmed.'}", styles["Normal"]),
    ]
    doc.build(story)
    return buffer.getvalue()


def generate_vendor_payments_list_excel(payments: List[VendorPayment], org: Optional[Organization] = None) -> bytes:
    """Export Vendor Payments made to Excel workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Payments Made"
    ws["A1"] = f"{(org.name if org else None) or 'Rooman Books'} - Payments Made to Vendors"
    ws["A2"] = f"Exported: {datetime.now().strftime('%d %b %Y, %I:%M %p')}"

    headers = ["Payment #", "Date", "Vendor", "Mode", "Reference", "Bill #", "Paid From", "Amount (INR)"]
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c_idx, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")

    for r_idx, p in enumerate(payments, start=5):
        ws.cell(row=r_idx, column=1, value=p.payment_number)
        ws.cell(row=r_idx, column=2, value=_fmt_date(p.date))
        ws.cell(row=r_idx, column=3, value=p.vendor.display_name if p.vendor else "-")
        ws.cell(row=r_idx, column=4, value=p.mode.upper())
        ws.cell(row=r_idx, column=5, value=p.reference or "-")
        ws.cell(row=r_idx, column=6, value=p.bill.bill_number if p.bill else "-")
        ws.cell(row=r_idx, column=7, value=p.bank_account.name if p.bank_account else "-")
        ws.cell(row=r_idx, column=8, value=float(p.amount or 0.0)).number_format = "₹#,##0.00"

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_vendor_payments_list_pdf(payments: List[VendorPayment], org: Optional[Organization] = None) -> bytes:
    """Generate Vendor Payments Made Registry PDF."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=24, leftMargin=24, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    th_style = ParagraphStyle("VPTH", fontName="Helvetica-Bold", fontSize=8, textColor=colors.white)
    td_style = ParagraphStyle("VPTD", fontName="Helvetica", fontSize=7.5, leading=9)
    td_right = ParagraphStyle("VPTDR", parent=td_style, alignment=2)

    story = [
        Paragraph(f"<b>{(org.name if org else None) or 'Rooman Books'} - Payments Made to Vendors</b>", styles["Heading2"]),
        Paragraph(f"Generated on {datetime.now().strftime('%d %b %Y, %I:%M %p')} · Total Payments: {len(payments)}", styles["Normal"]),
        Spacer(1, 12),
    ]

    rows = [[
        Paragraph("Voucher #", th_style),
        Paragraph("Vendor", th_style),
        Paragraph("Date", th_style),
        Paragraph("Mode", th_style),
        Paragraph("Reference", th_style),
        Paragraph("Bill #", th_style),
        Paragraph("Amount", th_style),
    ]]
    for p in payments:
        rows.append([
            Paragraph(p.payment_number, td_style),
            Paragraph((p.vendor.display_name if p.vendor else "-")[:24], td_style),
            Paragraph(_fmt_date(p.date), td_style),
            Paragraph(p.mode.upper(), td_style),
            Paragraph(p.reference or "-", td_style),
            Paragraph(p.bill.bill_number if p.bill else "Vendor Advance", td_style),
            Paragraph(_fmt_curr(p.amount), td_right),
        ])

    t = Table(rows, colWidths=[80, 130, 60, 60, 85, 75, 70])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563eb")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
    doc.build(story)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Expenses PDF / Excel
# ---------------------------------------------------------------------------


def generate_expenses_list_pdf(expenses: List[Expense], org: Optional[Organization] = None) -> bytes:
    """Generate an Expense Registry PDF report."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=24, leftMargin=24, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    th_style = ParagraphStyle("ETH", fontName="Helvetica-Bold", fontSize=8, textColor=colors.white)
    td_style = ParagraphStyle("ETD", fontName="Helvetica", fontSize=7.5, leading=9)
    td_right = ParagraphStyle("ETDR", parent=td_style, alignment=2)

    story = [
        Paragraph(f"<b>{(org.name if org else None) or 'Rooman Books'} - Operating Expenses Summary</b>", styles["Heading2"]),
        Paragraph(f"Generated on {datetime.now().strftime('%d %b %Y, %I:%M %p')} · Total Records: {len(expenses)}", styles["Normal"]),
        Spacer(1, 12),
    ]

    rows = [[
        Paragraph("Date", th_style),
        Paragraph("Category / Account", th_style),
        Paragraph("Payee", th_style),
        Paragraph("Payment Method", th_style),
        Paragraph("Reference", th_style),
        Paragraph("Amount", th_style),
    ]]
    for e in expenses:
        payee_str = e.vendor.display_name if e.vendor else (e.customer.display_name if e.customer else getattr(e, "payee", None) or "-")
        rows.append([
            Paragraph(_fmt_date(e.date), td_style),
            Paragraph((e.account.name if e.account else "-")[:24], td_style),
            Paragraph(payee_str[:20], td_style),
            Paragraph(e.payment_method.upper() if e.payment_method else "-", td_style),
            Paragraph(e.reference or "-", td_style),
            Paragraph(_fmt_curr(e.amount), td_right),
        ])

    t = Table(rows, colWidths=[65, 140, 110, 85, 80, 80])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
    doc.build(story)
    return buffer.getvalue()


def generate_expenses_list_excel(expenses: List[Expense], org: Optional[Organization] = None) -> bytes:
    """Export Operating Expenses to Excel workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Operating Expenses"
    ws["A1"] = f"{(org.name if org else None) or 'Rooman Books'} - Operating Expenses Journal"
    ws["A2"] = f"Exported: {datetime.now().strftime('%d %b %Y, %I:%M %p')}"

    headers = ["Date", "Expense Account", "Payee", "Payment Method", "Reference", "Description", "Amount (INR)"]
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c_idx, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")

    for r_idx, e in enumerate(expenses, start=5):
        payee_str = e.vendor.display_name if e.vendor else (e.customer.display_name if e.customer else getattr(e, "payee", None) or "-")
        ws.cell(row=r_idx, column=1, value=_fmt_date(e.date))
        ws.cell(row=r_idx, column=2, value=e.account.name if e.account else "-")
        ws.cell(row=r_idx, column=3, value=payee_str)
        ws.cell(row=r_idx, column=4, value=e.payment_method.upper() if e.payment_method else "-")
        ws.cell(row=r_idx, column=5, value=e.reference or "-")
        ws.cell(row=r_idx, column=6, value=e.notes or "-")
        ws.cell(row=r_idx, column=7, value=float(e.amount or 0.0)).number_format = "₹#,##0.00"

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Contacts (Customers & Vendors) PDF / Excel
# ---------------------------------------------------------------------------


def generate_contacts_list_pdf(contacts: List[Contact], contact_type: str = "customer", org: Optional[Organization] = None) -> bytes:
    """Generate Contacts / Customer Directory PDF."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=24, leftMargin=24, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    th_style = ParagraphStyle("CTH", fontName="Helvetica-Bold", fontSize=8, textColor=colors.white)
    td_style = ParagraphStyle("CTD", fontName="Helvetica", fontSize=7.5, leading=9)

    label = "Customer" if contact_type == "customer" else "Vendor"

    story = [
        Paragraph(f"<b>{(org.name if org else None) or 'Rooman Books'} - {label} Directory</b>", styles["Heading2"]),
        Paragraph(f"Total {label}s: {len(contacts)} · Generated on {datetime.now().strftime('%d %b %Y')}", styles["Normal"]),
        Spacer(1, 12),
    ]

    rows = [[
        Paragraph(f"{label} Name", th_style),
        Paragraph("Company / Legal", th_style),
        Paragraph("Email Address", th_style),
        Paragraph("Phone", th_style),
        Paragraph("GSTIN", th_style),
    ]]
    for c in contacts:
        rows.append([
            Paragraph(c.display_name, td_style),
            Paragraph(c.company_name or "-", td_style),
            Paragraph(c.email or "-", td_style),
            Paragraph(c.phone or "-", td_style),
            Paragraph(c.gstin or "-", td_style),
        ])

    t = Table(rows, colWidths=[130, 120, 130, 85, 95])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)
    doc.build(story)
    return buffer.getvalue()


def generate_contacts_list_excel(contacts: List[Contact], contact_type: str = "customer", org: Optional[Organization] = None) -> bytes:
    """Export Contacts / Customers to Excel workbook."""
    wb = Workbook()
    ws = wb.active
    label = "Customers" if contact_type == "customer" else "Vendors"
    ws.title = label
    ws["A1"] = f"{(org.name if org else None) or 'Rooman Books'} - {label} Master File"
    ws["A2"] = f"Exported: {datetime.now().strftime('%d %b %Y, %I:%M %p')}"

    headers = [f"{label[:-1]} Name", "Company Name", "Email Address", "Phone", "GSTIN", "PAN", "Billing Address"]
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c_idx, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    for r_idx, c in enumerate(contacts, start=5):
        ws.cell(row=r_idx, column=1, value=c.display_name)
        ws.cell(row=r_idx, column=2, value=c.company_name or "-")
        ws.cell(row=r_idx, column=3, value=c.email or "-")
        ws.cell(row=r_idx, column=4, value=c.phone or "-")
        ws.cell(row=r_idx, column=5, value=c.gstin or "-")
        ws.cell(row=r_idx, column=6, value=c.pan or "-")
        ws.cell(row=r_idx, column=7, value=c.billing_address or "-")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_overall_dashboard_report_pdf(summary_data: any, org: Optional[Organization] = None, period_label: str = "This Fiscal Year") -> bytes:
    """Generate a comprehensive Executive Financial & Operations Overall Report PDF."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0f172a"),
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#2563eb"),
    )
    sec_heading = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1e293b"),
    )
    value_style = ParagraphStyle(
        "ValueStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#334155"),
    )
    header_th = ParagraphStyle(
        "THStyle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
    )
    cell_style = ParagraphStyle(
        "CellStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#1e293b"),
    )
    cell_bold = ParagraphStyle(
        "CellBold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#0f172a"),
    )
    cell_right = ParagraphStyle(
        "CellRight",
        parent=cell_style,
        alignment=2,
    )
    cell_right_bold = ParagraphStyle(
        "CellRightBold",
        parent=cell_bold,
        alignment=2,
    )

    story = []

    org_name = (org.name if org else None) or "Rooman Technologies Pvt Ltd"
    org_addr = (org.address if org else None) or "Rooman House, #12 Rajajinagar"
    org_city_state = ", ".join(filter(None, [org.city if org else "Bengaluru", org.state if org else "Karnataka", org.postal_code if org else "560010"]))
    org_gstin = (org.gstin if org else None) or "29AABCR1234F1Z5"
    org_email = (org.email if org else None) or "shalya@rooman.com"

    header_table_data = [
        [
            Paragraph(f"<b>{org_name}</b>", title_style),
            Paragraph("<b>OVERALL FINANCIAL REPORT</b>", subtitle_style),
        ],
        [
            Paragraph(f"{org_addr}<br/>{org_city_state}<br/>GSTIN: {org_gstin}<br/>Email: {org_email}", value_style),
            Paragraph(
                f"<b>Period:</b> {period_label}<br/>"
                f"<b>Generated:</b> {datetime.now().strftime('%d %b %Y, %I:%M %p')}<br/>"
                f"<b>Scope:</b> Live Executive Dashboard Position",
                value_style,
            ),
        ],
    ]
    header_table = Table(header_table_data, colWidths=[330, 210])
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#2563eb"), spaceAfter=14))

    # Helper to read attr or dict
    def _g(obj, attr, default=None):
        if hasattr(obj, attr):
            return getattr(obj, attr)
        if isinstance(obj, dict):
            return obj.get(attr, default)
        return default

    # Key Performance Indicators Table
    story.append(Paragraph("Executive Financial Summary", sec_heading))
    story.append(Spacer(1, 6))

    rec = _g(summary_data, "receivables")
    pay = _g(summary_data, "payables")
    ie = _g(summary_data, "income_expense")
    inv = _g(summary_data, "inventory")

    tot_cash = float(_g(summary_data, "total_cash") or 0.0)
    tot_rec = float(_g(rec, "total_receivables") or 0.0)
    overdue_rec = float(_g(rec, "overdue_amount") or 0.0)
    tot_pay = float(_g(pay, "total_payables") or 0.0)
    overdue_pay = float(_g(pay, "overdue_amount") or 0.0)
    tot_inc = float(_g(ie, "total_income") or 0.0)
    tot_exp = float(_g(ie, "total_expense") or 0.0)
    net_val = float(_g(ie, "net") or 0.0)
    inv_val = float(_g(inv, "total_inventory_valuation") or 0.0)

    kpi_data = [
        [
            Paragraph("Metric", header_th),
            Paragraph("Amount (INR)", ParagraphStyle("THRight", parent=header_th, alignment=2)),
            Paragraph("Context / Breakdown", header_th),
        ],
        [
            Paragraph("Cash on Hand", cell_bold),
            Paragraph(_fmt_curr(tot_cash), cell_right_bold),
            Paragraph("Active liquid funds across bank and cash accounts", cell_style),
        ],
        [
            Paragraph("Accounts Receivable", cell_bold),
            Paragraph(_fmt_curr(tot_rec), cell_right_bold),
            Paragraph(f"Overdue: {_fmt_curr(overdue_rec)} ({_g(rec, 'total_unpaid_invoices', 0)} open invoices)", cell_style),
        ],
        [
            Paragraph("Accounts Payable", cell_bold),
            Paragraph(_fmt_curr(tot_pay), cell_right_bold),
            Paragraph(f"Overdue: {_fmt_curr(overdue_pay)} ({_g(pay, 'total_unpaid_bills', 0)} open bills)", cell_style),
        ],
        [
            Paragraph("Total Revenue / Income", cell_bold),
            Paragraph(_fmt_curr(tot_inc), cell_right_bold),
            Paragraph("Operating & sales revenues recorded in period", cell_style),
        ],
        [
            Paragraph("Total Operating Expenses", cell_bold),
            Paragraph(_fmt_curr(tot_exp), cell_right_bold),
            Paragraph("Expenses and direct operating costs", cell_style),
        ],
        [
            Paragraph("<b>Net Position (Profit / Loss)</b>", cell_bold),
            Paragraph(f"<b>{_fmt_curr(net_val)}</b>", cell_right_bold),
            Paragraph("Net operational surplus for selected period" if net_val >= 0 else "Net operational deficit for selected period", cell_style),
        ],
        [
            Paragraph("Inventory Valuation", cell_bold),
            Paragraph(_fmt_curr(inv_val), cell_right_bold),
            Paragraph(f"{_g(inv, 'total_items_count', 0)} items ({_g(inv, 'low_stock_items_count', 0)} low stock)", cell_style),
        ],
    ]

    kpi_table = Table(kpi_data, colWidths=[160, 130, 250])
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 14))

    # Bank Accounts Section
    bank_balances = _g(summary_data, "bank_balances") or []
    if bank_balances:
        story.append(Paragraph("Bank & Cash Accounts Position", sec_heading))
        story.append(Spacer(1, 5))
        bank_data = [
            [
                Paragraph("Account Name", header_th),
                Paragraph("Account Type", header_th),
                Paragraph("Current Balance", ParagraphStyle("THRight", parent=header_th, alignment=2)),
            ]
        ]
        for b in bank_balances:
            b_name = _g(b, "name") or "Account"
            b_type = (_g(b, "type") or "bank").replace("_", " ").title()
            b_bal = float(_g(b, "balance") or 0.0)
            bank_data.append([
                Paragraph(b_name, cell_style),
                Paragraph(b_type, cell_style),
                Paragraph(_fmt_curr(b_bal), cell_right),
            ])
        bank_table = Table(bank_data, colWidths=[240, 150, 150])
        bank_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(bank_table)
        story.append(Spacer(1, 14))

    # Top Customers Section
    top_custs = _g(summary_data, "top_customers") or []
    if top_custs:
        story.append(Paragraph("Top Customers by Sales Volume", sec_heading))
        story.append(Spacer(1, 5))
        cust_data = [
            [
                Paragraph("Customer Name", header_th),
                Paragraph("Total Billed Volume", ParagraphStyle("THRight", parent=header_th, alignment=2)),
            ]
        ]
        for tc in top_custs:
            c_name = _g(tc, "contact_name") or "Customer"
            c_amt = float(_g(tc, "amount") or 0.0)
            cust_data.append([
                Paragraph(c_name, cell_style),
                Paragraph(_fmt_curr(c_amt), cell_right),
            ])
        cust_table = Table(cust_data, colWidths=[360, 180])
        cust_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(cust_table)
        story.append(Spacer(1, 14))

    # Footer note
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceAfter=8))
    footer_text = f"Generated by Rooman Books Cloud Accounting &bull; {org_name} &bull; Confidential Management Report"
    story.append(Paragraph(footer_text, ParagraphStyle("FooterStyle", parent=value_style, alignment=1, textColor=colors.HexColor("#64748b"))))

    doc.build(story)
    return buffer.getvalue()


